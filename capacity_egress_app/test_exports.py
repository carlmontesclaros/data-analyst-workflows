import re
import pytest
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableFormula
from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters
import exports
import test_egress as t

def result():
    rooms = pd.DataFrame([t.room('301', 60.55, 30, 'site count'), t.room('302', 60.55)])
    for col, val in [('building_number', '999'), ('space_category', None), ('note', None),
                     ('fmis_capacity', None), ('in_latest_report', 1)]:
        rooms[col] = val
    rooms['visited'] = [1, 0]
    return t.egress.calculate(rooms, pd.DataFrame(t.TWO_DOORS, columns=t.EXIT_COLUMNS), t.factors())

def test_results_workbook_has_4_formatted_sheets(tmp_path):
    path = tmp_path / 'results.xlsx'
    counts = exports.write_results_workbook(result(), path)
    # S1 / S2 doors with no stair width -> 2 stair_not_measured warnings
    assert counts == {'floor_summary': 1, 'exit_detail': 2, 'room_detail': 2, 'warnings': 2}
    wb = load_workbook(path)
    assert wb.sheetnames == ['floor_summary', 'exit_detail', 'room_detail', 'warnings']
    ws = wb['floor_summary']
    assert ws.freeze_panes == 'A2' and ws['A1'].font.bold
    assert ws['A2'].fill.fgColor.rgb.endswith('C6EFCE')
    # results say which code edition they were screened under
    assert pd.read_excel(path, sheet_name='floor_summary')['code_edition'].tolist() == ['BCBC 2024']

# fake template like carl_space_report: BCBC factor table on top, header row 7, Table1 A-L, Pictures link formula
TEMPLATE_HEADER = ['Property', 'Floor', 'Space', 'Building Number', 'Net Space (sq m)', 'Space Category',
                   'Space Sub-Category', 'Capacity (Occupants)', 'Capacity (per code)',
                   'Room type (per code m^2 used in capacity)', 'Note', 'Pictures']
PICTURES = 'HYPERLINK("S:\\buildings\\"&SUBSTITUTE(A{r},"/","-")&"\\Floor "&B{r}&"\\"&C{r},"Open")'

def fake_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws['H1'], ws['H2'] = 'BCBC 2024 (per Table 3.1.17.1)', 'room type'
    for r, (rt, m2) in enumerate([('classroom', 1.85), ('lab', 4.6), ('non-fixed', 0.95), ('office', 9.3)], 3):
        ws.cell(r, 8, rt)
        ws.cell(r, 9, m2)
    for c, name in enumerate(TEMPLATE_HEADER, 1):
        ws.cell(7, c, name)
    # 301 is in the template with an FMIS count + note; 302 has a count but isn't visited; 999 the app doesn't have
    rows = [['Fake Hall', 3, 301, 999, 60.55, None, None, 0, 2230, None, 'by window', None],
            ['Fake Hall', 3, 302, 999, 60.55, None, None, 12, 2230, None, None, None],
            ['Fake Hall', 3, 999, 999, 1, None, None, None, None, None, 'old room', None]]
    for r, row in enumerate(rows, 8):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        ws.cell(r, 12, '=' + PICTURES.format(r=r))
    table = Table(displayName='Table1', ref='A7:L10')
    table.autoFilter = AutoFilter(ref='A7:L10')
    table.autoFilter.filterColumn.append(FilterColumn(colId=0, filters=Filters(filter=['Somewhere Else'])))
    table._initialise_columns()
    for column, name in zip(table.tableColumns, TEMPLATE_HEADER):
        column.name = name
    table.tableColumns[11].calculatedColumnFormula = TableFormula(attr_text=PICTURES.format(r=8))
    ws.add_table(table)
    wb.save(path)

def test_space_report_fills_the_template(tmp_path):
    template, path = tmp_path / 'template.xlsx', tmp_path / 'space_report.xlsx'
    fake_template(template)
    info = exports.write_space_report(result(), path, template)
    assert info['rooms'] == 2 and info['matched_template'] == 2
    df = pd.read_excel(path, header=6)
    assert list(df.columns) == TEMPLATE_HEADER and len(df) == 2
    r301, r302 = df.iloc[0], df.iloc[1]
    # site count for visited rooms; a room not visited keeps the template's count
    assert r301['Capacity (Occupants)'] == 30 and r302['Capacity (Occupants)'] == 12
    # per code = the app's load (60.55 m2 classroom -> 33), not the template's broken 2230
    assert r301['Capacity (per code)'] == 33 and r302['Capacity (per code)'] == 33
    # what the app doesn't store comes from the template's row
    assert r301['Note'] == 'by window'
    wb = load_workbook(path)
    ws = wb['Sheet1']
    assert ws['H3'].value == 'classroom' and ws['I6'].value == 9.3
    # Pictures link points at its own row
    assert ws['L8'].value == '=' + PICTURES.format(r=8) and ws['L9'].value == '=' + PICTURES.format(r=9)
    # table covers the new rows, old filter gone
    assert ws.tables['Table1'].ref == 'A7:L9' and not ws.tables['Table1'].autoFilter.filterColumn

def test_space_report_one_floor_and_no_template(tmp_path):
    template = tmp_path / 'template.xlsx'
    fake_template(template)
    info = exports.write_space_report(result(), tmp_path / 'none.xlsx', template, ['Nowhere'])
    # an empty selection still writes a valid table (one blank row)
    assert info['rooms'] == 0
    assert load_workbook(tmp_path / 'none.xlsx')['Sheet1'].tables['Table1'].ref == 'A7:L8'
    with pytest.raises(FileNotFoundError):
        exports.find_template(tmp_path / 'empty_folder')

def test_formula_for_row_leaves_text_alone():
    f = 'HYPERLINK("S:\\A8 photos\\"&amp;A8&amp;$B$8&amp;Table1[[#This Row],[Space]],"Open")'
    assert exports.formula_for_row(f, 8, 15) == 'HYPERLINK("S:\\A8 photos\\"&amp;A15&amp;$B$8&amp;Table1[[#This Row],[Space]],"Open")'

# capacity check summary .docx
def test_check_summary_docx(tmp_path):
    import zipfile
    from xml.dom.minidom import parseString
    rooms = pd.DataFrame([t.room('301', 60.55, 30, 'site count')])
    rooms['room_id'] = [7]
    check = t.egress.check_capacities(rooms, pd.DataFrame(t.TWO_DOORS, columns=t.EXIT_COLUMNS), t.factors(), {7: 40})
    when = pd.Timestamp('2026-10-13 10:30')
    name = exports.check_file_name(check, when)
    assert name == 'capacity_check_FH_301_2026-10-13.docx'
    path = exports.write_check_summary(check, tmp_path / name, 'Carl & co', when)
    with zipfile.ZipFile(path) as z:
        assert sorted(z.namelist()) == ['[Content_Types].xml', '_rels/.rels', 'word/_rels/document.xml.rels',
                                        'word/document.xml', 'word/styles.xml']
        doc = z.read('word/document.xml').decode('utf-8')
        for part in z.namelist():
            parseString(z.read(part))
    text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', doc))
    for expected in ['Fake Hall', 'Carl &amp; co', 'needs justification, +7 over per code', '30 (site count)',
                     'within exit capacity', 'Screening only, not an approval']:
        assert expected in text
    assert 'w:fill="FFEB9C"' in doc          # amber room row

def test_results_workbook_zone_colours(tmp_path):
    path = tmp_path / 'results.xlsx'
    exports.write_results_workbook(result(), path)
    assert load_workbook(path)['floor_summary']['A2'].fill.fgColor.rgb.endswith('C6EFCE')

def test_check_summary_shows_room_type_tried(tmp_path):
    import zipfile
    rooms = pd.DataFrame([t.room('301', 60.55)])
    rooms['room_id'] = [7]
    check = t.egress.check_capacities(rooms, pd.DataFrame(t.TWO_DOORS, columns=t.EXIT_COLUMNS), t.factors(), {}, {7: 'lab'})
    path = exports.write_check_summary(check, tmp_path / 'c.docx', 'Carl', pd.Timestamp('2026-10-16 10:00'))
    with zipfile.ZipFile(path) as z:
        text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', z.read('word/document.xml').decode('utf-8')))
    for expected in ['classroom -&gt; lab', '33 -&gt; 14', 'uses the per code load']:
        assert expected in text

def test_space_report_pm_change_on_a_visited_room_goes_into_capacity_occupants(tmp_path):
    # Carl, 2026-10-16: Capacity (Occupants) is the room capacity FMIS gets, so a PM change after the visit goes in
    rooms = pd.DataFrame([t.room('301', 60.55, 36, 'PM change'), t.room('302', 60.55)])
    for col, val in [('building_number', '999'), ('space_category', None), ('note', None),
                     ('fmis_capacity', None), ('in_latest_report', 1)]:
        rooms[col] = val
    rooms['visited'] = [1, 0]
    res = t.egress.calculate(rooms, pd.DataFrame(t.TWO_DOORS, columns=t.EXIT_COLUMNS), t.factors())
    template, path = tmp_path / 'template.xlsx', tmp_path / 'space_report.xlsx'
    fake_template(template)
    exports.write_space_report(res, path, template)
    assert pd.read_excel(path, header=6).iloc[0]['Capacity (Occupants)'] == 36
