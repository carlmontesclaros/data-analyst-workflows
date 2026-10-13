import os
import re
import glob
import zipfile
from collections import Counter
from xml.sax.saxutils import escape, unescape
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string
import capacity_db as db
import egress

# results workbook -> same 4 sheets as exit_capacity.py, column names people already know
SUMMARY_COLS = {'property': 'Property', 'floor': 'Floor', 'zone_wing': 'zone_wing', 'status': 'status',
                'flags': 'flags', 'total_load': 'total_load', 'exit_capacity': 'exit_capacity',
                'rooms_in_scope': 'rooms_in_scope', 'rooms_counted': 'rooms_counted',
                'counted_capacity': 'counted_capacity', 'area_based_load': 'area_based_load',
                'link_inflow': 'link_inflow', 'exit_count': 'exit_count', 'code_edition': 'code_edition',
                'floor_key': 'floor_key'}
EXIT_COLS = {'property': 'Property', 'floor': 'Floor', 'wing': 'wing', 'exit_id': 'exit_id',
             'exit_type': 'exit_type', 'clear_width_cm': 'clear_width_cm', 'counts_as_exit': 'counts_as_exit',
             'persons': 'persons', 'width_check': 'width_check', 'mm_per_person': 'mm_per_person',
             'minimum_mm': 'minimum_mm', 'minimum_mm_low_rise': 'minimum_mm_low_rise', 'into_wing': 'into_wing',
             'measured_date': 'measured_date', 'measured_by': 'measured_by', 'notes': 'notes', 'floor_key': 'floor_key'}
ROOM_COLS = {'property': 'Property', 'floor': 'Floor', 'space': 'Space', 'room_wing': 'room_wing',
             'zone_wing': 'zone_wing', 'sub_category': 'Space Sub-Category', 'room_type': 'room_type',
             'room_type_source': 'room_type_source', 'area_m2': 'Net Space (sq m)',
             'area_per_person_m2': 'area_per_person_m2', 'per_code_load': 'per_code_load',
             'capacity': 'room_capacity', 'capacity_source': 'capacity_source', 'fmis_capacity': 'fmis_capacity',
             'in_scope': 'in_scope', 'load_used': 'load_used', 'load_source': 'load_source',
             'in_latest_report': 'in_latest_report', 'floor_key': 'floor_key'}
WARNING_COLS = {'type': 'type', 'property': 'Property', 'floor': 'Floor', 'detail': 'detail'}

# row colours -> same as the app: red over exit capacity, amber other review, green within, grey not surveyed
COLOUR_FILLS = {egress.GREEN: PatternFill('solid', fgColor='C6EFCE'), egress.AMBER: PatternFill('solid', fgColor='FFEB9C'),
                egress.RED: PatternFill('solid', fgColor='FFC7CE'), egress.GREY: PatternFill('solid', fgColor='E7E6E6')}

def pick(df, properties=None, floors=None):
    if properties:
        df = df[df['property'].isin(properties)]
    if floors:
        df = df[df['floor'].isin(floors)]
    return df

def sheet(df, cols):
    return df[[c for c in cols if c in df.columns]].rename(columns=cols)

# formatting -> bold frozen header, filters, column widths <= 50
def format_sheets(writer):
    for ws in writer.sheets.values():
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col_cells in ws.columns:
            values = [len(str(c.value)) for c in col_cells if c.value is not None]
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(values, default=8) + 2, 50)

def write_results_workbook(result, path, properties=None, floors=None):
    summary = sheet(pick(result['zones'], properties, floors), SUMMARY_COLS)
    tables = {'floor_summary': summary,
              'exit_detail': sheet(pick(result['exits'], properties, floors), EXIT_COLS),
              'room_detail': sheet(pick(result['rooms'], properties, floors), ROOM_COLS),
              'warnings': sheet(pick(result['warnings'], properties, floors), WARNING_COLS)}
    try:
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            for name, df in tables.items():
                df.to_excel(writer, sheet_name=name, index=False)
            format_sheets(writer)
            # zone colours -> whole row on floor_summary
            ws = writer.sheets['floor_summary']
            colours = pick(result['zones'], properties, floors).apply(egress.zone_colour, axis=1).tolist() if len(summary) else []
            for row, colour in zip(ws.iter_rows(min_row=2), colours):
                for cell in row:
                    cell.fill = COLOUR_FILLS[colour]
    except PermissionError:
        raise PermissionError(f"can't write {path} - it's probably open in Excel. close it and try again")
    return {name: len(df) for name, df in tables.items()}

# space report export -> Carl's space report (carl_space_report), filled with the app's rooms
# it has slicers that openpyxl would destroy, so we never open it with openpyxl:
# copy the zip, rewrite only the room table's rows, keep every other file as is
TEMPLATE_DIR = os.path.join(db.DATA_DIR, "template")

# template columns, by header name. the app fills these; the rest come from the template's own row for that room
APP_COLS = {'Property': 'property', 'Floor': 'floor', 'Space': 'space', 'Building Number': 'building_number',
            'Net Space (sq m)': 'area_m2', 'Space Category': 'space_category', 'Space Sub-Category': 'sub_category'}
COUNT_COL, PER_CODE_COL = 'Capacity (Occupants)', 'Capacity (per code)'
ROOM_TYPE_COL, NOTE_COL = 'Room type (per code m^2 used in capacity)', 'Note'

def find_template(folder=TEMPLATE_DIR):
    # whatever .xlsx is in data/template/ -> no hardcoded filename
    files = sorted(f for f in glob.glob(os.path.join(folder, '*.xlsx')) if not os.path.basename(f).startswith('~$'))
    if not files:
        raise FileNotFoundError(f"no template in {folder} - put the space report template (.xlsx) there")
    return files[0]

def blank(value):
    return value is None or (isinstance(value, float) and pd.isna(value)) or (isinstance(value, str) and value.strip() == '')

def key(value):
    # same normalising as the import -> '3.0' and 3 both become '3'
    return '' if blank(value) else str(value).strip().removesuffix('.0')

def cell(ref, style, value):
    s = f' s="{style}"' if style else ''
    if blank(value):
        return f'<c r="{ref}"{s}/>'
    if isinstance(value, (int, float, np.integer, np.floating)) and not isinstance(value, bool):
        # pandas reads a column with blanks as float -> fund 10000.0 goes back as 10000
        return f'<c r="{ref}"{s}><v>{int(value) if float(value).is_integer() else value}</v></c>'
    text = str(value)
    # floor / space / building number like the FMIS export: 245 is a number, 'B303' or '0014' stays text
    if text.isdigit() and (text == '0' or not text.startswith('0')):
        return f'<c r="{ref}"{s}><v>{int(text)}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{escape(text)}</t></is></c>'

def template_parts(z):
    # the room table = the one with Property / Floor / Space in its header -> (table xml path, sheet xml path)
    table_path = None
    for name in z.namelist():
        if name.startswith('xl/tables/') and name.endswith('.xml'):
            names = [unescape(n) for n in re.findall(r'<tableColumn [^>]*?name="([^"]*)"', z.read(name).decode('utf-8'))]
            if {'Property', 'Floor', 'Space'} <= set(names):
                table_path = name
                break
    if table_path is None:
        raise ValueError("template has no table with Property, Floor and Space columns")
    # Excel writes Target="../tables/table1.xml", openpyxl writes Target="/xl/tables/table1.xml"
    target = re.compile(r'Target="[^"]*tables/' + re.escape(os.path.basename(table_path)) + '"')
    for name in z.namelist():
        if name.startswith('xl/worksheets/_rels/') and target.search(z.read(name).decode('utf-8')):
            return table_path, 'xl/worksheets/' + os.path.basename(name).removesuffix('.rels')
    raise ValueError(f"can't find the sheet for {table_path}")

def table_columns(table_xml):
    # column names in order + the calculated column formula of each (None if it has none), as written in the xml
    columns = re.findall(r'<tableColumn [^>]*?name="([^"]*)"[^>]*?(?:/>|>(.*?)</tableColumn>)', table_xml, re.S)
    names = [unescape(name) for name, _ in columns]
    formulas = [(re.search(r'<calculatedColumnFormula[^>]*>(.*?)</calculatedColumnFormula>', inner or '', re.S) or [None, None])[1]
                for _, inner in columns]
    return names, formulas

def formula_for_row(formula, first, r):
    # the table's formula points at its first data row (A8) -> point it at row r; text in quotes stays as is
    parts = formula.split('&quot;') if '&quot;' in formula else formula.split('"')
    quote = '&quot;' if '&quot;' in formula else '"'
    parts = [re.sub(rf'(?<![A-Za-z0-9_$])(\$?[A-Z]{{1,3}}){first}(?!\d)', rf'\g<1>{r}', p) if i % 2 == 0 else p
             for i, p in enumerate(parts)]
    return quote.join(parts)

def sheet_name(z, sheet_path):
    # xl/worksheets/sheet1.xml -> its tab name, via workbook.xml.rels (r:id) and workbook.xml (name)
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    rel_id = re.search(r'<Relationship [^>]*Id="(\w+)"[^>]*Target="[^"]*' + re.escape(os.path.basename(sheet_path)) + '"', rels)
    if rel_id is None:
        rel_id = re.search(r'<Relationship [^>]*Target="[^"]*' + re.escape(os.path.basename(sheet_path)) + r'"[^>]*Id="(\w+)"', rels)
    sheet = re.search(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="' + rel_id.group(1) + '"', z.read('xl/workbook.xml').decode('utf-8'))
    return unescape(sheet.group(1))

def row_styles(data_rows, n):
    # most common style per column, from the template's first 50 data rows
    styles = [Counter() for _ in range(n)]
    for row in data_rows[:50]:
        for c in re.finditer(r'<c r="([A-Z]+)\d+"(?: s="(\d+)")?', row):
            col = column_index_from_string(c.group(1)) - 1
            if col < n:
                styles[col][c.group(2)] += 1
    return [s.most_common(1)[0][0] if s else None for s in styles]

def write_space_report(result, path, template_path=None, properties=None, floors=None):
    template_path = template_path or find_template()
    rooms = pick(result['rooms'], properties, floors)
    rooms = rooms[rooms['in_latest_report'] == 1].sort_values(['property', 'floor', 'space'])

    with zipfile.ZipFile(template_path) as z:
        table_path, sheet_path = template_parts(z)
        table_xml = z.read(table_path).decode('utf-8')
        sheet_xml = z.read(sheet_path).decode('utf-8')
        header_row = int(re.search(r'<table [^>]*ref="[A-Z]+(\d+):', table_xml).group(1))
        names, formulas = table_columns(table_xml)
        n = len(names)
        col = {name: i for i, name in enumerate(names)}

        # the template's own rows -> baseline for what the app doesn't store (count of rooms not visited, notes, ...)
        baseline = pd.read_excel(template_path, sheet_name=sheet_name(z, sheet_path), header=header_row - 1)
        baseline_rows = {}
        for values in baseline.iloc[:, :n].itertuples(index=False):
            baseline_rows.setdefault((key(values[0]), key(values[1]), key(values[2])), list(values))

        start, end = sheet_xml.index('<sheetData>') + len('<sheetData>'), sheet_xml.index('</sheetData>')
        all_rows = re.findall(r'<row [^>]*/>|<row [^>]*>.*?</row>', sheet_xml[start:end], re.S)
        row_number = lambda row: int(re.search(r' r="(\d+)"', row).group(1))
        top_rows = [row for row in all_rows if row_number(row) <= header_row]
        styles = row_styles([row for row in all_rows if row_number(row) > header_row], n)
        style = lambda c: f' s="{styles[c]}"' if styles[c] else ''

        # new rows -> app values over the template's row for the same room
        first = header_row + 1
        new_rows = []
        matched = 0
        for i, room in enumerate(rooms.itertuples(index=False)):
            r = first + i
            k = (key(room.property), key(room.floor), key(room.space))
            values = list(baseline_rows.get(k, [None] * n))
            matched += k in baseline_rows
            for name, field in APP_COLS.items():
                if name in col:
                    values[col[name]] = getattr(room, field)
            # site count for visited rooms; the per code load the app works out (the report's own column is broken)
            if COUNT_COL in col and room.visited == 1 and not blank(room.capacity):
                values[col[COUNT_COL]] = int(room.capacity)
            if PER_CODE_COL in col:
                values[col[PER_CODE_COL]] = int(room.per_code_load)
            if ROOM_TYPE_COL in col and not blank(room.room_type_override):
                values[col[ROOM_TYPE_COL]] = room.room_type_override
            if NOTE_COL in col and not blank(room.note):
                values[col[NOTE_COL]] = room.note
            cells = [cell(f"{get_column_letter(c + 1)}{r}", styles[c], v) for c, v in enumerate(values)]
            # calculated columns (Pictures link) -> the table's formula for this row, Excel works it out on open
            for c, f in enumerate(formulas):
                if f:
                    cells[c] = f'<c r="{get_column_letter(c + 1)}{r}"{style(c)} t="str"><f>{formula_for_row(f, first, r)}</f></c>'
            new_rows.append(f'<row r="{r}" spans="1:{n}">' + ''.join(cells) + '</row>')
        if not new_rows:
            # an Excel table needs at least one data row
            new_rows.append(f'<row r="{first}" spans="1:{n}">'
                            + ''.join(cell(f"{get_column_letter(c + 1)}{first}", styles[c], None) for c in range(n)) + '</row>')
        last = first + len(new_rows) - 1

        sheet_xml = sheet_xml[:start] + ''.join(top_rows) + ''.join(new_rows) + sheet_xml[end:]
        sheet_xml = re.sub(r'<dimension ref="([A-Z]+)1:([A-Z]+)\d+"/>', rf'<dimension ref="\g<1>1:\g<2>{last}"/>', sheet_xml)
        sheet_xml = re.sub(r'topLeftCell="[A-Z]+\d+"', f'topLeftCell="A{first}"', sheet_xml)
        sheet_xml = re.sub(r'activeCell="[A-Z]+\d+" sqref="[^"]*"', f'activeCell="A{first}" sqref="A{first}"', sheet_xml)
        sheet_xml = re.sub(r'(<dataValidation [^>]*sqref=")([A-Z]+)\d+:([A-Z]+)\d+"', rf'\g<1>\g<2>{first}:\g<3>{last}"', sheet_xml)

        # table covers the new rows; clear the saved filter (the template was saved filtered)
        table_xml = re.sub(r'ref="([A-Z]+)(\d+):([A-Z]+)\d+"', rf'ref="\g<1>\g<2>:\g<3>{last}"', table_xml)
        table_xml = re.sub(r'<filterColumn .*?</filterColumn>', '', table_xml, flags=re.S)

        try:
            with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as out:
                for item in z.infolist():
                    name = item.filename
                    if name == 'xl/calcChain.xml':
                        continue        # old list of formula cells; Excel rebuilds it
                    data = z.read(name)
                    if name == sheet_path:
                        data = sheet_xml.encode('utf-8')
                    elif name == table_path:
                        data = table_xml.encode('utf-8')
                    elif name.startswith('xl/pivotCache/pivotCacheDefinition') and b'refreshOnLoad' not in data:
                        # pivots still hold the template's numbers -> refresh when the file opens
                        data = data.replace(b'<pivotCacheDefinition ', b'<pivotCacheDefinition refreshOnLoad="1" ', 1)
                    elif name == 'xl/workbook.xml' and b'fullCalcOnLoad' not in data:
                        data = data.replace(b'<calcPr ', b'<calcPr fullCalcOnLoad="1" ', 1)
                    elif name == 'xl/_rels/workbook.xml.rels':
                        data = re.sub(rb'<Relationship [^>]*Target="calcChain.xml"/>', b'', data)
                    elif name == '[Content_Types].xml':
                        data = re.sub(rb'<Override PartName="/xl/calcChain.xml"[^>]*/>', b'', data)
                    out.writestr(item, data)
        except PermissionError:
            raise PermissionError(f"can't write {path} - it's probably open in Excel. close it and try again")
    return {'rooms': len(rooms), 'matched_template': int(matched), 'template': os.path.basename(template_path)}


# capacity check summary -> one-page .docx to attach to the email to Mark
# a .docx is a zip of xml files, so it's written here directly: nothing extra to install
FILL = {egress.GREEN: 'C6EFCE', egress.AMBER: 'FFEB9C', egress.RED: 'FFC7CE', egress.GREY: 'E7E6E6', egress.NONE: None}
CHECK_FOOTER = ("Screening only, not an approval. An editor reviews every change before it is saved. "
                "Email this summary to Mark Arrojado with the reason for the request.")

DOCX_TYPES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
              '<Default Extension="xml" ContentType="application/xml"/>'
              '<Override PartName="/word/document.xml" '
              'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
              '<Override PartName="/word/styles.xml" '
              'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>')
DOCX_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
             '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
             'Target="word/document.xml"/></Relationships>')
DOCX_DOC_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
                 'Target="styles.xml"/></Relationships>')
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
DOCX_STYLES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               f'<w:styles {W}><w:docDefaults><w:rPrDefault><w:rPr>'
               '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/><w:sz w:val="20"/></w:rPr></w:rPrDefault>'
               '<w:pPrDefault><w:pPr><w:spacing w:after="80"/></w:pPr></w:pPrDefault></w:docDefaults></w:styles>')

def para(text, bold=False, size=None, colour=None, after=None):
    rpr = ('<w:b/>' if bold else '') + (f'<w:color w:val="{colour}"/>' if colour else '') + \
          (f'<w:sz w:val="{size * 2}"/>' if size else '')
    ppr = f'<w:pPr><w:spacing w:after="{after}"/></w:pPr>' if after is not None else ''
    return (f'<w:p>{ppr}<w:r>{"<w:rPr>" + rpr + "</w:rPr>" if rpr else ""}'
            f'<w:t xml:space="preserve">{escape(str(text))}</w:t></w:r></w:p>')

def table(header, rows, fills, widths):
    # widths in twentieths of a point; fills = one colour per row (or None)
    border = ''.join(f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
                     for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV'])
    def row(values, fill, bold=False):
        cells = ''
        for value, width in zip(values, widths):
            shade = f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>' if fill else ''
            cells += (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{shade}</w:tcPr>'
                      f'{para(value, bold=bold, after=0)}</w:tc>')
        return f'<w:tr>{cells}</w:tr>'
    grid = ''.join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    return (f'<w:tbl><w:tblPr><w:tblW w:w="{sum(widths)}" w:type="dxa"/><w:tblBorders>{border}</w:tblBorders>'
            f'<w:tblCellMar><w:left w:w="60" w:type="dxa"/><w:right w:w="60" w:type="dxa"/></w:tblCellMar></w:tblPr>'
            f'<w:tblGrid>{grid}</w:tblGrid>{row(header, "D9D9D9", bold=True)}'
            + ''.join(row(r, f) for r, f in zip(rows, fills)) + '</w:tbl>')

def check_file_name(check, when=None):
    first = check['rooms'][0]
    initials = ''.join(word[0] for word in re.findall(r'[A-Za-z]+', first['property']) if word[0].isupper())
    rooms = '_'.join(sanitize_part(r['space']) for r in check['rooms'][:3]) + ('_more' if len(check['rooms']) > 3 else '')
    return f"capacity_check_{initials}_{rooms}_{(when or pd.Timestamp.now()):%Y-%m-%d}.docx"

def sanitize_part(text):
    return re.sub(r'[^A-Za-z0-9-]+', '-', str(text)).strip('-')

def write_check_summary(check, path, user_name, when=None):
    when = when or pd.Timestamp.now()
    if not check['rooms']:
        raise ValueError("the check has no rooms")
    prop = check['rooms'][0]['property']
    body = [para("Room capacity check", bold=True, size=16, after=40),
            para(f"{prop}   |   {when:%B %d, %Y %H:%M}   |   run by {user_name}   |   {check['code_edition']}",
                 colour='595959', after=200)]

    body.append(para("Rooms", bold=True, size=12, after=60))
    rows, fills = [], []
    for r in check['rooms']:
        current = '-' if r['current'] is None else r['current']
        scope = '' if r['in_scope'] else ' (out of scope: not counted toward the floor)'
        rows.append([f"Floor {r['floor']}  {r['space']}", r['room_type'] or '-', f"{current} ({r['current_source']})",
                     r['proposed'], r['per_code_load'], r['verdict'] + scope])
        fills.append(FILL[r['colour']])
    body.append(table(['Room', 'Type', 'Capacity now', 'Proposed', 'Per code load', 'Result'], rows, fills,
                      [1500, 1100, 1500, 1000, 1100, 3400]))

    body.append(para("", after=120))
    body.append(para("Floor exits: occupant load vs exit capacity", bold=True, size=12, after=60))
    rows, fills = [], []
    for z in check['zones']:
        zone = f"Floor {z['floor']}" + (f" wing {z['zone_wing']}" if z['zone_wing'] else '')
        capacity = '-' if z['verdict'] == egress.CANT_CHECK else z['exit_capacity']
        result = z['verdict'] + (f" (also: {z['other_flags']})" if z['other_flags'] else '')
        rows.append([zone + (' *' if z['has_checked_room'] else ''), f"{z['load_before']} -> {z['load_after']}", capacity, result])
        fills.append(FILL[z['colour']])
    body.append(table(['Zone', 'Occupant load now -> with change', 'Exit capacity', 'Result'], rows, fills,
                      [1800, 2400, 1300, 4100]))
    body.append(para("* zone holding a room in this check", colour='595959', after=200))

    body.append(para("What the results mean", bold=True, after=40))
    for line in ["Within per code: the proposed number is at or under the building code's area-based number for the room.",
                 "Needs justification: over the per code load. Allowed only when the higher number can be shown "
                 "(fixed seats, a posted sign; BCBC 2024 3.1.17.1.(1)(c)), so an editor looks into it.",
                 "Over exit capacity: the floor's measured exits can't carry the occupant load with this change.",
                 egress.CANT_CHECK]:
        body.append(para("- " + line, after=20))
    body.append(para("", after=120))
    body.append(para(CHECK_FOOTER, bold=True))

    sect = ('<w:sectPr><w:pgSz w:w="12240" w:h="15840"/>'
            '<w:pgMar w:top="1000" w:right="1000" w:bottom="1000" w:left="1000" w:header="500" w:footer="500" w:gutter="0"/>'
            '</w:sectPr>')
    document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>'
                + ''.join(body) + sect + '</w:body></w:document>')
    try:
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as out:
            out.writestr('[Content_Types].xml', DOCX_TYPES)
            out.writestr('_rels/.rels', DOCX_RELS)
            out.writestr('word/_rels/document.xml.rels', DOCX_DOC_RELS)
            out.writestr('word/styles.xml', DOCX_STYLES)
            out.writestr('word/document.xml', document)
    except PermissionError:
        raise PermissionError(f"can't write {path} - it's probably open in Word. close it and try again")
    return path
